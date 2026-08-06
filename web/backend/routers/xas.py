"""XAS / XANES API endpoints.

Handles TEY/TFY dual-channel XAS data from beamline DAT files.
"""

from __future__ import annotations

import io
import json
import re
import zipfile
from typing import Any, Dict, List, Optional

import numpy as np
import pandas as pd
from fastapi import APIRouter, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from scipy.signal import savgol_filter

from core.parsers import looks_like_excel, numeric_excel_table
from core.peak_fitting import fit_peaks, perturb_init_peaks
from core.processing import apply_normalization
from core.spectrum_ops import interpolate_spectrum_to_grid, mean_spectrum_arrays
from core.xas_global_fitting import fit_xas_global_components
from db.xas_database import get_sample_edge_peaks, list_samples

router = APIRouter()


# ── parsing helpers (self-contained, no streamlit dependency) ─────────────────

def _is_numeric_line(line: str) -> bool:
    parts = line.strip().replace(",", " ").split()
    if not parts:
        return False
    try:
        for part in parts:
            float(part)
        return True
    except ValueError:
        return False


def _extract_column_names(header_line: "str | None", n_cols: int) -> list[str]:
    """Try to parse a header line into column names for n_cols columns."""
    if header_line:
        for sep in ("\t", ",", r"\s+"):
            try:
                parts_df = pd.read_csv(io.StringIO(header_line), sep=sep, header=None, engine="python")
                parts = [str(p).strip() for p in parts_df.iloc[0].tolist() if str(p).strip()]
                if len(parts) >= n_cols:
                    return parts[:n_cols]
            except Exception:
                pass
        # simple whitespace split fallback
        parts = header_line.split()
        if len(parts) >= n_cols:
            return parts[:n_cols]
    return [f"Col {i + 1}" for i in range(n_cols)]


def _auto_detect_xas_column_mapping(column_names: list[str]) -> dict[str, "int | None"]:
    """Guess energy/tey/tfy/io column indices from column names."""
    lower = [n.lower().replace("-", "").replace("_", "").replace(" ", "") for n in column_names]

    def find(keywords: list[str]) -> "int | None":
        for kw in keywords:
            for i, n in enumerate(lower):
                if kw in n:
                    return i
        return None

    energy_idx = find(["energy", "ev", "mono", "hv", "energyev", "ener"])
    io_idx = find(["i0", "io", "incident", "ring", "mesh", "flux", "curmd02", "curmd2", "ringcur"])
    tey_idx = find(["tey", "drain", "totalelectron", "pey", "curmd01", "curmd1"])
    tfy_idx = find(["tfy", "fluorescence", "pfy", "fy", "curmd03", "curmd3"])

    n = len(column_names)
    if energy_idx is None:
        energy_idx = 0

    if n >= 6 and tey_idx is None and tfy_idx is None and io_idx is None:
        # 6-col beamline heuristic: Energy, Phase, Gap, TFY, TEY, I0
        tfy_idx, tey_idx, io_idx = 3, 4, 5

    if tey_idx is None or tfy_idx is None:
        used = {energy_idx}
        if io_idx is not None:
            used.add(io_idx)
        remaining = [i for i in range(n) if i not in used]
        if tey_idx is None and remaining:
            tey_idx = remaining.pop(0)
        if tfy_idx is None and remaining:
            tfy_idx = remaining.pop(0)

    return {
        "energy": int(energy_idx),
        "tey": int(tey_idx) if tey_idx is not None else 1,
        "tfy": int(tfy_idx) if tfy_idx is not None else 2,
        "io": int(io_idx) if io_idx is not None else None,
    }


def _parse_xas_table_bytes(raw: bytes):
    """Parse text-like XAS/DAT files.

    Returns (df_with_int_columns, column_names, error).
    df rows are NaN-dropped but NOT sorted; column order matches column_names.
    """
    excel_df, excel_err = numeric_excel_table(raw, min_columns=3)
    if excel_df is not None:
        col_names = [str(c) for c in excel_df.columns]
        excel_df.columns = list(range(excel_df.shape[1]))
        return excel_df, col_names, None
    if excel_err and looks_like_excel(raw):
        return None, [], excel_err

    for enc in ("utf-8", "utf-8-sig", "big5", "cp950", "latin-1", "utf-16"):
        try:
            text = raw.decode(enc)
        except UnicodeDecodeError:
            continue

        lines = text.splitlines()

        # Find data block start and capture potential header line
        header_candidate: "str | None" = None
        data_start = -1
        for i, line in enumerate(lines):
            stripped = line.strip()
            if not stripped:
                continue
            if _is_numeric_line(line):
                data_start = i
                break
            elif stripped[0] not in ("#", "%", ";", "!"):
                header_candidate = stripped  # last non-comment, non-numeric line

        if data_start >= 0:
            numeric_lines: list[str] = []
            for line in lines[data_start:]:
                if _is_numeric_line(line):
                    numeric_lines.append(line.strip())
                else:
                    break
            if len(numeric_lines) >= 2:
                clean = "\n".join(numeric_lines)
                for sep in ("\t", ",", r"\s+"):
                    try:
                        df = pd.read_csv(io.StringIO(clean), sep=sep, header=None, engine="python")
                        num = df.apply(pd.to_numeric, errors="coerce")
                        valid = [col for col in num.columns if num[col].notna().mean() > 0.8]
                        if len(valid) >= 3:
                            out = num[valid].dropna(how="any").copy()
                            n_cols = out.shape[1]
                            out.columns = list(range(n_cols))
                            col_names = _extract_column_names(header_candidate, n_cols)
                            return out.reset_index(drop=True), col_names, None
                    except Exception:
                        pass

        clean_lines = [
            line for line in lines
            if line.strip() and line.strip()[0] not in ("#", "%", ";", "!")
        ]
        if clean_lines:
            clean = "\n".join(clean_lines)
            for sep in (",", "\t", r"\s+"):
                for header in (0, None):
                    try:
                        df = pd.read_csv(io.StringIO(clean), sep=sep, header=header, engine="python")
                        num = df.apply(pd.to_numeric, errors="coerce")
                        valid = [col for col in num.columns if num[col].notna().mean() > 0.8]
                        if len(valid) >= 3:
                            out = num[valid].dropna(how="any").copy()
                            n_cols = out.shape[1]
                            if header == 0:
                                col_names = [str(df.columns[c]) for c in valid]
                            else:
                                col_names = _extract_column_names(header_candidate, n_cols)
                            out.columns = list(range(n_cols))
                            return out.reset_index(drop=True), col_names, None
                    except Exception:
                        pass

    return None, [], "無法解析：請確認檔案至少包含 Energy、TEY、TFY 三欄數字"


def _apply_column_mapping(
    df: pd.DataFrame,
    col_map: dict[str, "int | None"],
    flip_tfy: bool,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, str | None]:
    """Apply column mapping to df → sorted (energy, tey, tfy) arrays."""
    n = df.shape[1]
    e_idx = col_map.get("energy", 0) or 0
    tey_idx = col_map.get("tey") or 1
    tfy_idx = col_map.get("tfy") or 2
    io_idx = col_map.get("io")

    if max(e_idx, tey_idx, tfy_idx) >= n:
        return np.array([]), np.array([]), np.array([]), "欄位索引超出範圍"

    energy = df.iloc[:, e_idx].to_numpy(dtype=float)
    tey_raw = df.iloc[:, tey_idx].to_numpy(dtype=float)
    tfy_raw = df.iloc[:, tfy_idx].to_numpy(dtype=float)

    if io_idx is not None and io_idx < n:
        io_arr = df.iloc[:, io_idx].to_numpy(dtype=float)
        denom = np.where(np.abs(io_arr) > 1e-30, np.abs(io_arr), np.nan)
        tey_raw = tey_raw / denom
        tfy_raw = tfy_raw / denom

    if flip_tfy:
        tfy_raw = 1.0 - tfy_raw

    mask = np.isfinite(energy) & np.isfinite(tey_raw) & np.isfinite(tfy_raw)
    if np.count_nonzero(mask) < 2:
        return energy, np.array([]), np.array([]), "有效資料點不足"

    energy = energy[mask]
    tey_raw = tey_raw[mask]
    tfy_raw = tfy_raw[mask]
    order = np.argsort(energy)
    return energy[order], tey_raw[order], tfy_raw[order], None


def _find_white_line(x: np.ndarray, y: np.ndarray, e_min: float, e_max: float) -> float | None:
    mask = (x >= e_min) & (x <= e_max)
    if not np.any(mask):
        return None
    idx = int(np.argmax(y[mask]))
    return float(x[mask][idx])


def _normalize_post_edge(
    x: np.ndarray, y: np.ndarray,
    edge_region: tuple[float, float],
    norm_region: tuple[float, float],
) -> tuple[np.ndarray, float]:
    """Normalize by post-edge step height using a linear pre/post baseline."""
    pre_mask = (x >= edge_region[0]) & (x <= edge_region[1])
    post_mask = (x >= norm_region[0]) & (x <= norm_region[1])

    if not np.any(pre_mask) or not np.any(post_mask):
        return y, 1.0

    pre_mean = float(np.mean(y[pre_mask]))
    post_mean = float(np.mean(y[post_mask]))
    edge_step = post_mean - pre_mean
    if abs(edge_step) < 1e-20:
        return y, 1.0

    normalized = (y - pre_mean) / edge_step
    return normalized, edge_step


def _detect_e0(x: np.ndarray, y: np.ndarray) -> float:
    """Detect edge energy E₀ as the energy of the maximum first derivative."""
    if len(x) < 3:
        return float(x[len(x) // 2])
    dy = np.gradient(y, x)
    return float(x[int(np.argmax(dy))])


def _normalize_athena(
    x: np.ndarray,
    y: np.ndarray,
    pre_region: tuple[float, float],
    post_region: tuple[float, float],
    e0: Optional[float] = None,
    pre_deg: int = 1,
    post_deg: int = 2,
) -> tuple[np.ndarray, np.ndarray, float, float, np.ndarray, np.ndarray, np.ndarray]:
    """
    Athena-style XANES normalization.

    Steps:
    1. Fit a line to the pre-edge region and subtract it from the whole spectrum.
    2. Detect (or use) E₀.
    3. Fit a polynomial to the post-edge region of the subtracted spectrum.
    4. edge_step = post_poly(E₀).
    5. normalized = y_sub / edge_step  (~0 before edge, ~1 after edge).
    6. flattened  = (y_sub − post_poly) / edge_step + 1  (post-edge goes flat at 1).

    Returns (normalized, flattened, e0_used, edge_step, pre_line, post_poly, y_sub).
    pre_line / post_poly / y_sub are retained for visualization purposes.
    """
    if e0 is None:
        e0 = _detect_e0(x, y)

    # --- pre-edge linear fit & subtraction ---
    pre_mask = (x >= pre_region[0]) & (x <= pre_region[1])
    n_pre = int(pre_mask.sum())
    if n_pre >= 2:
        pre_coeffs = np.polyfit(x[pre_mask], y[pre_mask], min(pre_deg, n_pre - 1))
    elif n_pre == 1:
        pre_coeffs = np.array([float(y[pre_mask][0])])
    else:
        pre_coeffs = np.array([float(np.mean(y))])
    pre_line = np.polyval(pre_coeffs, x)
    y_sub = y - pre_line

    # --- post-edge polynomial fit ---
    post_mask = (x >= post_region[0]) & (x <= post_region[1])
    n_post = int(post_mask.sum())
    if n_post >= 2:
        post_coeffs = np.polyfit(x[post_mask], y_sub[post_mask], min(post_deg, n_post - 1))
    elif n_post == 1:
        post_coeffs = np.array([float(y_sub[post_mask][0])])
    else:
        post_coeffs = np.array([float(np.mean(y_sub))])
    post_poly = np.polyval(post_coeffs, x)

    # --- edge step at E₀ ---
    edge_step = float(np.polyval(post_coeffs, e0))
    if abs(edge_step) < 1e-20:
        return y, y, float(e0), 1.0, pre_line, post_poly, y_sub

    normalized = y_sub / edge_step
    flattened = (y_sub - post_poly) / edge_step + 1.0

    return normalized, flattened, float(e0), float(edge_step), pre_line, post_poly, y_sub


def _exafs_window(kind: str, n: int) -> np.ndarray:
    if n <= 1:
        return np.ones(n, dtype=float)
    if kind == "none":
        return np.ones(n, dtype=float)
    return np.hanning(n)


def _calculate_exafs_preview(
    energy: np.ndarray,
    mu: np.ndarray,
    e0: float,
    k_min: float,
    k_max: float,
    k_weight: int,
    rbkg: float,
    window: str,
    r_max: float,
    edge_step: Optional[float] = None,
) -> tuple[dict[str, Any], list[str]]:
    """Build a lightweight EXAFS preview.

    This is intentionally a preview backend, not a replacement for Larch autobk.
    It keeps the API boundary ready for a future autobk/xftf implementation.
    """
    warnings: list[str] = []
    if energy.size != mu.size or energy.size < 8:
        raise ValueError("EXAFS 輸入資料點不足")
    if not np.isfinite(e0):
        raise ValueError("E0 必須是有限數值")

    finite = np.isfinite(energy) & np.isfinite(mu)
    energy = energy[finite]
    mu = mu[finite]
    if energy.size < 8:
        raise ValueError("EXAFS 有效資料點不足")
    order = np.argsort(energy)
    energy = energy[order]
    mu = mu[order]

    post_mask = energy > float(e0)
    if int(np.count_nonzero(post_mask)) < 16:
        raise ValueError("E0 之後的 post-edge 資料點不足，無法建立 χ(k)")

    k_raw = np.sqrt(np.maximum(energy[post_mask] - float(e0), 0.0) / 3.80998212)
    mu_raw = mu[post_mask]
    k_lo = float(min(k_min, k_max))
    k_hi = float(max(k_min, k_max))
    k_mask = np.isfinite(k_raw) & np.isfinite(mu_raw) & (k_raw >= k_lo) & (k_raw <= k_hi)
    if int(np.count_nonzero(k_mask)) < 16:
        raise ValueError("目前 k range 內資料點不足，請調整 k min / k max 或 E0")

    k_src = k_raw[k_mask]
    mu_src = mu_raw[k_mask]
    src_order = np.argsort(k_src)
    k_src = k_src[src_order]
    mu_src = mu_src[src_order]

    target_count = int(np.clip(round((float(k_src[-1]) - float(k_src[0])) / 0.035), 160, 900))
    k_grid = np.linspace(float(k_src[0]), float(k_src[-1]), target_count)
    dk = float(k_grid[1] - k_grid[0]) if target_count > 1 else 1.0
    mu_grid = np.interp(k_grid, k_src, mu_src)

    smooth_points = int(np.clip(round((max(float(rbkg), 0.2) * 10.0) / max(dk, 1e-9)), 5, max(5, target_count // 3)))
    if smooth_points % 2 == 0:
      smooth_points += 1
    if smooth_points >= target_count:
      smooth_points = target_count - 1 if (target_count - 1) % 2 == 1 else target_count - 2
    smooth_points = max(5, smooth_points)

    try:
        poly_order = min(3, smooth_points - 2)
        mu0 = savgol_filter(mu_grid, smooth_points, poly_order, mode="interp")
    except Exception:
        warnings.append("Savitzky-Golay μ0 平滑失敗，已改用移動平均")
        kernel = np.ones(smooth_points, dtype=float) / float(smooth_points)
        mu0 = np.convolve(mu_grid, kernel, mode="same")

    if edge_step is not None and np.isfinite(edge_step) and abs(edge_step) > 1e-12:
        step = float(abs(edge_step))
    else:
        span = float(np.nanmax(mu_grid) - np.nanmin(mu_grid))
        step = max(abs(span), 1e-9)
        warnings.append("未提供 edge step，已用輸入 μ(E) span 做預覽尺度")

    chi = (mu_grid - mu0) / step
    k_weight_int = int(np.clip(round(k_weight), 0, 3))
    chi_weighted = chi * np.power(k_grid, k_weight_int)
    win = _exafs_window(window, target_count)
    r_values = np.linspace(0.0, max(float(r_max), 1.0), int(round(max(float(r_max), 1.0) / 0.02)) + 1)
    ft_re: list[float] = []
    ft_im: list[float] = []
    for r in r_values:
        phase = 2.0 * k_grid * float(r)
        amp = chi_weighted * win
        ft_re.append(float(np.sum(amp * np.cos(phase)) * dk))
        ft_im.append(float(np.sum(amp * np.sin(phase)) * dk))
    ft_re_arr = np.array(ft_re, dtype=float)
    ft_im_arr = np.array(ft_im, dtype=float)
    ft_mag = np.hypot(ft_re_arr, ft_im_arr)

    log = [
        "EXAFS backend preview: SciPy Savitzky-Golay μ0 smoothing + direct FT",
        f"E0: {float(e0):.6g} eV",
        f"k range: {k_lo:.6g}-{k_hi:.6g} A^-1",
        f"k-weight: {k_weight_int}",
        f"Rbkg preview: {float(rbkg):.6g} A",
        f"FT window: {window}",
        f"R max: {float(r_max):.6g} A",
        f"smooth points: {smooth_points}",
        "Note: this is a visual preview, not a full Larch autobk/xftf replacement.",
    ]

    return {
        "energy": (float(e0) + 3.80998212 * np.square(k_grid)).tolist(),
        "mu": mu_grid.tolist(),
        "mu0": mu0.tolist(),
        "k": k_grid.tolist(),
        "chi": chi.tolist(),
        "chi_weighted": chi_weighted.tolist(),
        "r": r_values.tolist(),
        "ft_mag": ft_mag.tolist(),
        "ft_re": ft_re_arr.tolist(),
        "ft_im": ft_im_arr.tolist(),
        "edge_step": float(step),
        "smooth_points": int(smooth_points),
        "method": "scipy_savgol_preview",
        "log": log,
    }, warnings


def _calculate_exafs_larch(
    energy: np.ndarray,
    mu: np.ndarray,
    e0: float,
    k_min: float,
    k_max: float,
    k_weight: int,
    rbkg: float,
    window: str,
    r_max: float,
    edge_step: Optional[float] = None,
) -> tuple[dict[str, Any], list[str]]:
    """Run EXAFS with Larch when available, otherwise fall back safely.

    The import is intentionally lazy so deployments without xraylarch keep
    working. Larch APIs have changed across versions, so this wrapper catches
    failures and returns the SciPy preview with an explicit warning.
    """
    try:
        from larch import Group  # type: ignore
        from larch.xafs import autobk, xftf  # type: ignore
    except Exception as exc:
        result, warnings = _calculate_exafs_preview(
            energy, mu, e0, k_min, k_max, k_weight, rbkg, window, r_max, edge_step
        )
        result["method"] = "scipy_savgol_preview_fallback"
        result["log"] = [
            "Larch requested but xraylarch is not available; falling back to SciPy preview.",
            *result.get("log", []),
        ]
        return result, [f"Larch 未安裝或無法載入，已 fallback 到 SciPy preview：{exc}", *warnings]

    try:
        finite = np.isfinite(energy) & np.isfinite(mu)
        energy = energy[finite]
        mu = mu[finite]
        order = np.argsort(energy)
        energy = energy[order]
        mu = mu[order]
        if energy.size < 16:
            raise ValueError("EXAFS 有效資料點不足")

        group = Group(energy=energy, mu=mu)
        k_lo = float(min(k_min, k_max))
        k_hi = float(max(k_min, k_max))
        k_weight_int = int(np.clip(round(k_weight), 0, 3))

        try:
            autobk(energy, mu, group=group, rbkg=float(rbkg), e0=float(e0), kmin=k_lo, kmax=k_hi, kweight=k_weight_int)
        except TypeError:
            autobk(energy, mu, group=group, rbkg=float(rbkg), e0=float(e0))

        k_arr = np.array(getattr(group, "k"), dtype=float)
        chi_arr = np.array(getattr(group, "chi"), dtype=float)
        mask = np.isfinite(k_arr) & np.isfinite(chi_arr) & (k_arr >= k_lo) & (k_arr <= k_hi)
        if int(np.count_nonzero(mask)) < 16:
            raise ValueError("Larch autobk 回傳的 k range 內資料點不足")
        k_arr = k_arr[mask]
        chi_arr = chi_arr[mask]

        energy_k = float(e0) + 3.80998212 * np.square(k_arr)
        mu_k = np.interp(energy_k, energy, mu)
        bkg = getattr(group, "bkg", None)
        if bkg is not None:
            mu0_k = np.interp(energy_k, energy, np.array(bkg, dtype=float))
        else:
            step = float(abs(edge_step)) if edge_step is not None and np.isfinite(edge_step) and abs(edge_step) > 1e-12 else 1.0
            mu0_k = mu_k - chi_arr * step
        chi_weighted = chi_arr * np.power(k_arr, k_weight_int)

        try:
            xftf(k_arr, chi_arr, group=group, kmin=k_lo, kmax=k_hi, kweight=k_weight_int, window=window, rmax=float(r_max))
        except TypeError:
            xftf(k_arr, chi_arr, group=group, kmin=k_lo, kmax=k_hi, kweight=k_weight_int)

        r_arr = np.array(getattr(group, "r"), dtype=float)
        ft_mag = np.array(getattr(group, "chir_mag"), dtype=float)
        ft_re = np.array(getattr(group, "chir_re"), dtype=float)
        ft_im = np.array(getattr(group, "chir_im"), dtype=float)
        r_mask = np.isfinite(r_arr) & (r_arr <= max(float(r_max), 1.0))
        r_arr = r_arr[r_mask]
        ft_mag = ft_mag[r_mask]
        ft_re = ft_re[r_mask]
        ft_im = ft_im[r_mask]

        log = [
            "EXAFS backend: Larch autobk() + xftf()",
            f"E0: {float(e0):.6g} eV",
            f"k range: {k_lo:.6g}-{k_hi:.6g} A^-1",
            f"k-weight: {k_weight_int}",
            f"Rbkg: {float(rbkg):.6g} A",
            f"FT window: {window}",
            f"R max: {float(r_max):.6g} A",
            "Note: R-space peak positions are not phase-corrected bond lengths.",
        ]
        return {
            "energy": energy_k.tolist(),
            "mu": mu_k.tolist(),
            "mu0": mu0_k.tolist(),
            "k": k_arr.tolist(),
            "chi": chi_arr.tolist(),
            "chi_weighted": chi_weighted.tolist(),
            "r": r_arr.tolist(),
            "ft_mag": ft_mag.tolist(),
            "ft_re": ft_re.tolist(),
            "ft_im": ft_im.tolist(),
            "edge_step": float(abs(edge_step)) if edge_step is not None and np.isfinite(edge_step) and abs(edge_step) > 1e-12 else 1.0,
            "smooth_points": 0,
            "method": "larch_autobk_xftf",
            "log": log,
        }, []
    except Exception as exc:
        result, warnings = _calculate_exafs_preview(
            energy, mu, e0, k_min, k_max, k_weight, rbkg, window, r_max, edge_step
        )
        result["method"] = "scipy_savgol_preview_fallback"
        result["log"] = [
            f"Larch calculation failed; falling back to SciPy preview: {exc}",
            *result.get("log", []),
        ]
        return result, [f"Larch 計算失敗，已 fallback 到 SciPy preview：{exc}", *warnings]


# ── Gaussian template helpers ─────────────────────────────────────────────────

_GAUSS_SIGMA_FACTOR = 2.3548200450309493  # 2√(2ln2)


def _gaussian(x: np.ndarray, center: float, fwhm: float, amplitude: float) -> np.ndarray:
    sigma = fwhm / _GAUSS_SIGMA_FACTOR
    return amplitude * np.exp(-0.5 * ((x - center) / sigma) ** 2)


def _fit_gaussian_center(
    x: np.ndarray, y: np.ndarray,
    center: float, fwhm: float, amplitude: float,
    search_range: float,
) -> float:
    """Return best center via dot-product cross-correlation within ±search_range."""
    if search_range <= 0:
        return center
    sigma = fwhm / _GAUSS_SIGMA_FACTOR
    c_lo = max(float(x[0]), center - search_range)
    c_hi = min(float(x[-1]), center + search_range)
    candidates = np.linspace(c_lo, c_hi, 200)
    scores = np.array([
        float(np.dot(y, amplitude * np.exp(-0.5 * ((x - c) / sigma) ** 2)))
        for c in candidates
    ])
    return float(candidates[int(np.argmax(scores))])


# ── pydantic models ───────────────────────────────────────────────────────────

class GaussPeak(BaseModel):
    center: float
    fwhm: float
    amplitude: float


class XasColumnMapping(BaseModel):
    energy: int = 0
    tey: int = 1
    tfy: int = 2
    io: Optional[int] = None


class ParsedXasFile(BaseModel):
    name: str
    x: List[float]
    tey: List[float]
    tfy: List[float]
    mapping: Dict[str, Any]
    n_cols: int
    column_names: List[str] = Field(default_factory=list)
    raw_columns: List[List[float]] = Field(default_factory=list)
    default_mapping: Optional[XasColumnMapping] = None


class ParseResponse(BaseModel):
    files: List[ParsedXasFile]
    errors: List[str] = Field(default_factory=list)


class DatasetInput(BaseModel):
    name: str
    x: List[float]
    tey: List[float]
    tfy: List[float]


class ProcessParams(BaseModel):
    interpolate: bool = False
    n_points: int = 2000
    average: bool = False
    energy_shift: float = 0.0
    norm_method: str = "none"         # none | min_max | max | area | post_edge | mean_region | athena_norm
    e0_override: Optional[float] = None   # manual E₀ for athena_norm; None = auto-detect
    norm_tey_start: Optional[float] = None
    norm_tey_end: Optional[float] = None
    norm_tfy_start: Optional[float] = None
    norm_tfy_end: Optional[float] = None
    norm_tey_pre_start: Optional[float] = None  # for post_edge: pre-edge region start
    norm_tey_pre_end: Optional[float] = None    # for post_edge: pre-edge region end
    norm_tfy_pre_start: Optional[float] = None
    norm_tfy_pre_end: Optional[float] = None
    white_line_start: Optional[float] = None
    white_line_end: Optional[float] = None
    gauss_enabled: bool = False
    gauss_channel: str = "both"             # both | TEY | TFY
    gauss_peaks: List[GaussPeak] = Field(default_factory=list)
    gauss_search: float = 0.5              # ±eV center search range
    d2y_enabled: bool = False


class ProcessRequest(BaseModel):
    datasets: List[DatasetInput]
    params: ProcessParams


class ProcessedDataset(BaseModel):
    name: str
    x: List[float]
    tey_raw: List[float]
    tfy_raw: List[float]
    tey_processed: List[float]
    tfy_processed: List[float]
    white_line_tey: Optional[float] = None
    white_line_tfy: Optional[float] = None
    edge_step_tey: Optional[float] = None
    edge_step_tfy: Optional[float] = None
    e0_tey: Optional[float] = None
    e0_tfy: Optional[float] = None
    tey_flattened: Optional[List[float]] = None
    tfy_flattened: Optional[List[float]] = None
    tey_gaussian: Optional[List[float]] = None
    tfy_gaussian: Optional[List[float]] = None
    tey_after_gauss: Optional[List[float]] = None
    tfy_after_gauss: Optional[List[float]] = None
    tey_d2y: Optional[List[float]] = None
    tfy_d2y: Optional[List[float]] = None
    tey_pre_edge_line: Optional[List[float]] = None
    tfy_pre_edge_line: Optional[List[float]] = None
    tey_post_edge_poly: Optional[List[float]] = None
    tfy_post_edge_poly: Optional[List[float]] = None
    tey_pre_subtracted: Optional[List[float]] = None
    tfy_pre_subtracted: Optional[List[float]] = None


class ProcessResponse(BaseModel):
    datasets: List[ProcessedDataset]
    average: Optional[ProcessedDataset] = None


class ExafsRequest(BaseModel):
    energy: List[float]
    mu: List[float]
    e0: float
    k_min: float = 2.0
    k_max: float = 10.0
    k_weight: int = 2
    rbkg: float = 1.0
    window: str = "hanning"
    r_max: float = 6.0
    edge_step: Optional[float] = None
    backend_method: str = "scipy_preview"  # scipy_preview | larch_autobk


class ExafsResponse(BaseModel):
    energy: List[float]
    mu: List[float]
    mu0: List[float]
    k: List[float]
    chi: List[float]
    chi_weighted: List[float]
    r: List[float]
    ft_mag: List[float]
    ft_re: List[float]
    ft_im: List[float]
    edge_step: float
    smooth_points: int
    method: str
    warnings: List[str] = Field(default_factory=list)
    log: List[str] = Field(default_factory=list)


# ── endpoints ─────────────────────────────────────────────────────────────────

@router.post("/parse", response_model=ParseResponse)
async def parse_xas_files(
    files: List[UploadFile] = File(...),
    flip_tfy: bool = True,
):
    results: list[ParsedXasFile] = []
    errors: list[str] = []

    for uf in files:
        raw = await uf.read()
        df, col_names, err = _parse_xas_table_bytes(raw)
        if err or df is None:
            errors.append(f"{uf.filename}: {err or '解析失敗'}")
            continue

        n_cols = int(df.shape[1])
        if not col_names:
            col_names = [f"Col {i + 1}" for i in range(n_cols)]

        # Auto-detect column mapping
        detected = _auto_detect_xas_column_mapping(col_names)
        col_map = XasColumnMapping(**detected)

        # Apply mapping to get sorted x/tey/tfy
        energy, tey_arr, tfy_arr, prep_err = _apply_column_mapping(df, detected, flip_tfy)
        if prep_err and len(tey_arr) < 2:
            errors.append(f"{uf.filename}: {prep_err}")
            continue

        # Build raw_columns (unsorted, for frontend re-mapping)
        raw_columns = [df.iloc[:, i].tolist() for i in range(n_cols)]

        results.append(ParsedXasFile(
            name=uf.filename or "unknown",
            x=energy.tolist(),
            tey=tey_arr.tolist(),
            tfy=tfy_arr.tolist(),
            mapping={"mode": "column_mapped", "flip_tfy": bool(flip_tfy)},
            n_cols=n_cols,
            column_names=col_names,
            raw_columns=raw_columns,
            default_mapping=col_map,
        ))

    return ParseResponse(files=results, errors=errors)


@router.post("/process", response_model=ProcessResponse)
def process_xas(req: ProcessRequest):
    p = req.params
    datasets = req.datasets

    if not datasets:
        raise HTTPException(status_code=400, detail="沒有資料集")

    processed_datasets: list[ProcessedDataset] = []

    for ds in datasets:
        x = np.array(ds.x, dtype=float)
        tey = np.array(ds.tey, dtype=float)
        tfy = np.array(ds.tfy, dtype=float)

        # energy shift
        x_shifted = x + p.energy_shift

        # interpolation
        if p.interpolate:
            x_grid = np.linspace(float(x_shifted.min()), float(x_shifted.max()), int(p.n_points))
            tey = np.interp(x_grid, x_shifted, tey)
            tfy = np.interp(x_grid, x_shifted, tfy)
            x_shifted = x_grid

        x_out = x_shifted

        tey_proc = tey.copy()
        tfy_proc = tfy.copy()

        # gaussian template subtraction (before BG)
        tey_gauss_model: np.ndarray | None = None
        tfy_gauss_model: np.ndarray | None = None
        tey_after_gauss: np.ndarray | None = None
        tfy_after_gauss: np.ndarray | None = None

        if p.gauss_enabled and p.gauss_peaks:
            apply_tey = p.gauss_channel in ("both", "TEY")
            apply_tfy = p.gauss_channel in ("both", "TFY")

            if apply_tey:
                model = np.zeros_like(tey_proc)
                for gp in p.gauss_peaks:
                    c = _fit_gaussian_center(x_out, tey_proc, gp.center, gp.fwhm, gp.amplitude, p.gauss_search)
                    model += _gaussian(x_out, c, gp.fwhm, gp.amplitude)
                tey_gauss_model = model
                tey_proc = tey_proc - model
                tey_after_gauss = tey_proc.copy()

            if apply_tfy:
                model = np.zeros_like(tfy_proc)
                for gp in p.gauss_peaks:
                    c = _fit_gaussian_center(x_out, tfy_proc, gp.center, gp.fwhm, gp.amplitude, p.gauss_search)
                    model += _gaussian(x_out, c, gp.fwhm, gp.amplitude)
                tfy_gauss_model = model
                tfy_proc = tfy_proc - model
                tfy_after_gauss = tfy_proc.copy()

        # normalization
        edge_step_tey: float | None = None
        edge_step_tfy: float | None = None
        e0_tey_val: float | None = None
        e0_tfy_val: float | None = None
        tey_flat: np.ndarray | None = None
        tfy_flat: np.ndarray | None = None
        tey_pre_line_arr: np.ndarray | None = None
        tfy_pre_line_arr: np.ndarray | None = None
        tey_post_poly_arr: np.ndarray | None = None
        tfy_post_poly_arr: np.ndarray | None = None
        tey_pre_sub_arr: np.ndarray | None = None
        tfy_pre_sub_arr: np.ndarray | None = None

        if p.norm_method == "athena_norm":
            x_min = float(np.min(x_out))
            x_max = float(np.max(x_out))
            x_span = max(x_max - x_min, 1e-12)
            tey_pre_start = p.norm_tey_pre_start if p.norm_tey_pre_start is not None else x_min
            tey_pre_end   = p.norm_tey_pre_end   if p.norm_tey_pre_end   is not None else x_min + x_span * 0.3
            tey_post_start = p.norm_tey_start if p.norm_tey_start is not None else x_min + x_span * 0.7
            tey_post_end   = p.norm_tey_end   if p.norm_tey_end   is not None else x_max
            tfy_pre_start = p.norm_tfy_pre_start if p.norm_tfy_pre_start is not None else x_min
            tfy_pre_end   = p.norm_tfy_pre_end   if p.norm_tfy_pre_end   is not None else x_min + x_span * 0.3
            tfy_post_start = p.norm_tfy_start if p.norm_tfy_start is not None else x_min + x_span * 0.7
            tfy_post_end   = p.norm_tfy_end   if p.norm_tfy_end   is not None else x_max
            tey_proc, tey_flat, e0_t, step_t, tey_pre_line_arr, tey_post_poly_arr, tey_pre_sub_arr = _normalize_athena(
                x_out, tey_proc,
                (tey_pre_start, tey_pre_end),
                (tey_post_start, tey_post_end),
                e0=p.e0_override,
            )
            tfy_proc, tfy_flat, e0_f, step_f, tfy_pre_line_arr, tfy_post_poly_arr, tfy_pre_sub_arr = _normalize_athena(
                x_out, tfy_proc,
                (tfy_pre_start, tfy_pre_end),
                (tfy_post_start, tfy_post_end),
                e0=p.e0_override,
            )
            edge_step_tey = float(step_t)
            edge_step_tfy = float(step_f)
            e0_tey_val = float(e0_t)
            e0_tfy_val = float(e0_f)
        elif p.norm_method == "post_edge":
            x_min = float(np.min(x_out))
            x_max = float(np.max(x_out))
            x_span = max(x_max - x_min, 1e-12)
            tey_pre_start = p.norm_tey_pre_start if p.norm_tey_pre_start is not None else x_min
            tey_pre_end = p.norm_tey_pre_end if p.norm_tey_pre_end is not None else x_min + x_span * 0.3
            tey_norm_start = p.norm_tey_start if p.norm_tey_start is not None else x_min + x_span * 0.7
            tey_norm_end = p.norm_tey_end if p.norm_tey_end is not None else x_max
            tfy_pre_start = p.norm_tfy_pre_start if p.norm_tfy_pre_start is not None else x_min
            tfy_pre_end = p.norm_tfy_pre_end if p.norm_tfy_pre_end is not None else x_min + x_span * 0.3
            tfy_norm_start = p.norm_tfy_start if p.norm_tfy_start is not None else x_min + x_span * 0.7
            tfy_norm_end = p.norm_tfy_end if p.norm_tfy_end is not None else x_max
            tey_proc, step_t = _normalize_post_edge(
                x_out,
                tey_proc,
                (tey_pre_start, tey_pre_end),
                (tey_norm_start, tey_norm_end),
            )
            tfy_proc, step_f = _normalize_post_edge(
                x_out,
                tfy_proc,
                (tfy_pre_start, tfy_pre_end),
                (tfy_norm_start, tfy_norm_end),
            )
            edge_step_tey = float(step_t)
            edge_step_tfy = float(step_f)
        elif p.norm_method != "none" and p.norm_method != "post_edge":
            tey_proc = apply_normalization(
                x_out,
                tey_proc,
                norm_method=p.norm_method,
                norm_x_start=p.norm_tey_start,
                norm_x_end=p.norm_tey_end,
            )
            tfy_proc = apply_normalization(
                x_out,
                tfy_proc,
                norm_method=p.norm_method,
                norm_x_start=p.norm_tfy_start,
                norm_x_end=p.norm_tfy_end,
            )

        # white line
        wl_tey: float | None = None
        wl_tfy: float | None = None
        if p.white_line_start is not None and p.white_line_end is not None:
            wl_tey = _find_white_line(x_out, tey_proc, p.white_line_start, p.white_line_end)
            wl_tfy = _find_white_line(x_out, tfy_proc, p.white_line_start, p.white_line_end)

        # second derivative
        tey_d2y: np.ndarray | None = None
        tfy_d2y: np.ndarray | None = None
        if p.d2y_enabled and len(x_out) > 4:
            tey_d2y = np.gradient(np.gradient(tey_proc, x_out), x_out)
            tfy_d2y = np.gradient(np.gradient(tfy_proc, x_out), x_out)

        processed_datasets.append(ProcessedDataset(
            name=ds.name,
            x=x_out.tolist(),
            tey_raw=tey.tolist(),
            tfy_raw=tfy.tolist(),
            tey_processed=tey_proc.tolist(),
            tfy_processed=tfy_proc.tolist(),
            white_line_tey=wl_tey,
            white_line_tfy=wl_tfy,
            edge_step_tey=edge_step_tey,
            edge_step_tfy=edge_step_tfy,
            e0_tey=e0_tey_val,
            e0_tfy=e0_tfy_val,
            tey_flattened=tey_flat.tolist() if tey_flat is not None else None,
            tfy_flattened=tfy_flat.tolist() if tfy_flat is not None else None,
            tey_gaussian=tey_gauss_model.tolist() if tey_gauss_model is not None else None,
            tfy_gaussian=tfy_gauss_model.tolist() if tfy_gauss_model is not None else None,
            tey_after_gauss=tey_after_gauss.tolist() if tey_after_gauss is not None else None,
            tfy_after_gauss=tfy_after_gauss.tolist() if tfy_after_gauss is not None else None,
            tey_d2y=tey_d2y.tolist() if tey_d2y is not None else None,
            tfy_d2y=tfy_d2y.tolist() if tfy_d2y is not None else None,
            tey_pre_edge_line=tey_pre_line_arr.tolist() if tey_pre_line_arr is not None else None,
            tfy_pre_edge_line=tfy_pre_line_arr.tolist() if tfy_pre_line_arr is not None else None,
            tey_post_edge_poly=tey_post_poly_arr.tolist() if tey_post_poly_arr is not None else None,
            tfy_post_edge_poly=tfy_post_poly_arr.tolist() if tfy_post_poly_arr is not None else None,
            tey_pre_subtracted=tey_pre_sub_arr.tolist() if tey_pre_sub_arr is not None else None,
            tfy_pre_subtracted=tfy_pre_sub_arr.tolist() if tfy_pre_sub_arr is not None else None,
        ))

    # average across all datasets
    average_ds: ProcessedDataset | None = None
    if p.average and len(processed_datasets) > 1:
        try:
            x_ref = np.array(processed_datasets[0].x)
            tey_arrays = [np.interp(x_ref, np.array(d.x), np.array(d.tey_processed)) for d in processed_datasets]
            tfy_arrays = [np.interp(x_ref, np.array(d.x), np.array(d.tfy_processed)) for d in processed_datasets]
            tey_avg = np.mean(tey_arrays, axis=0)
            tfy_avg = np.mean(tfy_arrays, axis=0)
            wl_tey_avg: float | None = None
            wl_tfy_avg: float | None = None
            if p.white_line_start is not None and p.white_line_end is not None:
                wl_tey_avg = _find_white_line(x_ref, tey_avg, p.white_line_start, p.white_line_end)
                wl_tfy_avg = _find_white_line(x_ref, tfy_avg, p.white_line_start, p.white_line_end)
            average_ds = ProcessedDataset(
                name="平均",
                x=x_ref.tolist(),
                tey_raw=tey_avg.tolist(),
                tfy_raw=tfy_avg.tolist(),
                tey_processed=tey_avg.tolist(),
                tfy_processed=tfy_avg.tolist(),
                white_line_tey=wl_tey_avg,
                white_line_tfy=wl_tfy_avg,
            )
        except Exception:
            pass

    return ProcessResponse(datasets=processed_datasets, average=average_ds)


@router.post("/exafs", response_model=ExafsResponse)
def process_exafs(req: ExafsRequest):
    try:
        args = dict(
            energy=np.array(req.energy, dtype=float),
            mu=np.array(req.mu, dtype=float),
            e0=float(req.e0),
            k_min=float(req.k_min),
            k_max=float(req.k_max),
            k_weight=int(req.k_weight),
            rbkg=float(req.rbkg),
            window=req.window if req.window in ("hanning", "none") else "hanning",
            r_max=float(req.r_max),
            edge_step=req.edge_step,
        )
        if req.backend_method == "larch_autobk":
            result, warnings = _calculate_exafs_larch(**args)
        else:
            result, warnings = _calculate_exafs_preview(**args)
        return ExafsResponse(**result, warnings=warnings)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"EXAFS 處理失敗：{exc}") from exc


# ── XANES deconvolution ───────────────────────────────────────────────────────

class DeconvPeak(BaseModel):
    center: float
    delta: float = 2.0
    name: str = ""
    ptype: str = "gaussian"   # gaussian | lorentzian


class DeconvRequest(BaseModel):
    x: List[float]
    y: List[float]
    peaks: List[DeconvPeak]
    fwhm_inst: float = 0.5
    fwhm_init: float = 1.5
    link_fwhm: bool = False
    include_step: bool = True
    e0: float = 0.0
    fit_lo: Optional[float] = None
    fit_hi: Optional[float] = None


class DeconvParamRow(BaseModel):
    name: str
    value: float
    stderr: float
    vary: bool


class DeconvResponse(BaseModel):
    success: bool
    x_fit: List[float]
    y_fit: List[float]
    components: Dict[str, List[float]]
    residual: List[float]
    r_factor: float
    params_table: List[DeconvParamRow]
    message: str = ""


@router.post("/deconv", response_model=DeconvResponse)
def xanes_deconv(req: DeconvRequest):
    try:
        from lmfit.models import GaussianModel, LorentzianModel, StepModel
        from lmfit import Parameters
    except ImportError:
        raise HTTPException(status_code=500, detail="lmfit 未安裝，無法執行 XANES 去卷積擬合")

    x = np.array(req.x, dtype=float)
    y = np.array(req.y, dtype=float)

    lo = req.fit_lo if req.fit_lo is not None else float(x[0])
    hi = req.fit_hi if req.fit_hi is not None else float(x[-1])
    mask = (x >= lo) & (x <= hi)
    if mask.sum() < 4:
        raise HTTPException(status_code=400, detail="擬合範圍內資料點不足")

    x_f = x[mask]
    y_f = y[mask]

    _SIGMA_FACTOR = 2.3548200450309493
    sigma_inst = req.fwhm_inst / _SIGMA_FACTOR
    sigma_init = req.fwhm_init / _SIGMA_FACTOR

    composite = None
    params = Parameters()

    if req.include_step:
        step = StepModel(form='arctan', prefix='step_')
        composite = step
        params.update(step.make_params(
            center=dict(value=req.e0, vary=False),
            amplitude=dict(value=1.0, vary=False),
            sigma=dict(value=sigma_inst, min=sigma_inst * 0.5),
        ))

    for i, pk in enumerate(req.peaks):
        prefix = f'p{i}_'
        comp = LorentzianModel(prefix=prefix) if pk.ptype == 'lorentzian' else GaussianModel(prefix=prefix)
        p = comp.make_params()
        p[f'{prefix}center'].set(value=pk.center, min=pk.center - pk.delta, max=pk.center + pk.delta)
        p[f'{prefix}amplitude'].set(value=0.3, min=0)
        if req.link_fwhm and req.include_step and i == 0:
            p[f'{prefix}sigma'].set(value=sigma_init, min=sigma_inst)
        elif req.link_fwhm and req.include_step and i > 0:
            p[f'{prefix}sigma'].set(expr='p0_sigma', min=sigma_inst)
        else:
            p[f'{prefix}sigma'].set(value=sigma_init, min=sigma_inst)
        params.update(p)
        composite = comp if composite is None else composite + comp

    if composite is None:
        raise HTTPException(status_code=400, detail="沒有任何擬合分量")

    try:
        fit_result = composite.fit(y_f, params, x=x_f, method='leastsq')
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"擬合失敗：{exc}") from exc

    y_fit_arr = np.array(fit_result.best_fit)
    residual = y_f - y_fit_arr
    ss_res = float(np.sum(residual ** 2))
    ss_tot = float(np.sum(y_f ** 2))
    r_factor = ss_res / ss_tot if ss_tot > 0 else 0.0

    comps_raw = fit_result.eval_components(x=x_f)
    components: Dict[str, List[float]] = {k: list(map(float, v)) for k, v in comps_raw.items()}

    param_rows = [
        DeconvParamRow(
            name=name,
            value=float(par.value),
            stderr=float(par.stderr) if par.stderr is not None else 0.0,
            vary=bool(par.vary),
        )
        for name, par in fit_result.params.items()
    ]

    return DeconvResponse(
        success=bool(fit_result.success),
        x_fit=x_f.tolist(),
        y_fit=y_fit_arr.tolist(),
        components=components,
        residual=residual.tolist(),
        r_factor=r_factor,
        params_table=param_rows,
        message=str(fit_result.message) if fit_result.message else "",
    )


# ── Peak fitting (XPS-compatible, adapted for XAS energy axis) ────────────────

class XasGlobalDataset(BaseModel):
    name: str
    x: List[float]
    y: List[float]


class XasGlobalSmallPeak(BaseModel):
    range: List[float] = Field(default_factory=lambda: [544.0, 545.9], min_length=2, max_length=2)
    center: float = 545.0
    center_min: float = 544.0
    center_max: float = 545.9
    lock_center: bool = False
    fwhm: float = 1.0
    fwhm_min: float = 0.2
    fwhm_max: float = 2.0
    lock_fwhm: bool = False
    background: str = "linear"


class XasGlobalMainPeak(BaseModel):
    label: str
    center: float
    center_min: float
    center_max: float
    lock_center: bool = False
    fwhm: float = 1.2
    fwhm_min: float = 0.6
    fwhm_max: float = 2.0
    lock_fwhm: bool = False


class XasGlobalFitRequest(BaseModel):
    datasets: List[XasGlobalDataset]
    small_peak: XasGlobalSmallPeak
    main_peaks: List[XasGlobalMainPeak]
    fit_range: List[float] = Field(min_length=2, max_length=2)
    main_background: str = "linear"
    ratio_numerator: Optional[str] = "C2"
    ratio_denominator: Optional[str] = "B2"
    max_nfev: int = 30000


class XasGlobalExportRequest(BaseModel):
    config: Dict[str, Any]
    result: Dict[str, Any]


def _global_fit_excel(config: dict, result: dict) -> bytes:
    buffer = io.BytesIO()
    shared_rows = result.get("shared_peaks", [])
    sample_rows = [{
        "sample": dataset.get("name"),
        "r_squared": dataset.get("r_squared"),
        "rmse": dataset.get("rmse"),
        "ratio_numerator": result.get("ratio_numerator"),
        "ratio_denominator": result.get("ratio_denominator"),
        "area_ratio": dataset.get("area_ratio"),
        "small_peak_center_eV": dataset.get("small_peak", {}).get("center"),
        "small_peak_fwhm_eV": dataset.get("small_peak", {}).get("fwhm"),
        "small_peak_area": dataset.get("small_peak", {}).get("area"),
    } for dataset in result.get("datasets", [])]
    peak_rows = []
    for dataset in result.get("datasets", []):
        for peak in dataset.get("peaks", []):
            peak_rows.append({"sample": dataset.get("name"), **peak})
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        pd.DataFrame(shared_rows).to_excel(writer, sheet_name="共用峰參數", index=False)
        pd.DataFrame(sample_rows).to_excel(writer, sheet_name="擬合品質與面積比", index=False)
        pd.DataFrame(peak_rows).to_excel(writer, sheet_name="各樣品峰結果", index=False)
        pd.DataFrame([{"config_json": json.dumps(config, ensure_ascii=False)}]).to_excel(writer, sheet_name="設定", index=False)
        for index, dataset in enumerate(result.get("datasets", []), start=1):
            components = dataset.get("components", [])
            peaks = dataset.get("peaks", [])
            columns: dict[str, Any] = {
                "Energy_eV": dataset.get("x", []),
                "Original": dataset.get("original", []),
                "Small_Gaussian": dataset.get("small_peak", {}).get("component", []),
                "Small_Local_Background": dataset.get("small_peak", {}).get("local_background", []),
                "Corrected": dataset.get("corrected", []),
                "Main_Background": dataset.get("background", []),
                "Total_Fit": dataset.get("total_fit", []),
                "Residual": dataset.get("residual", []),
            }
            for peak_index, component in enumerate(components):
                label = peaks[peak_index].get("label", f"Peak_{peak_index + 1}") if peak_index < len(peaks) else f"Peak_{peak_index + 1}"
                columns[f"Peak_{label}"] = component
            safe_name = re.sub(r"[\\/*?:\[\]]", "_", str(dataset.get("name") or f"sample_{index}"))[:22]
            pd.DataFrame(columns).to_excel(writer, sheet_name=f"{index:02d}_{safe_name}"[:31], index=False)
    return buffer.getvalue()


@router.post("/global-fit")
def fit_xas_global(req: XasGlobalFitRequest):
    if req.small_peak.background not in {"constant", "linear"}:
        raise HTTPException(status_code=400, detail="小峰背景僅支援 constant 或 linear")
    if req.main_background not in {"constant", "linear"}:
        raise HTTPException(status_code=400, detail="主擬合背景僅支援 constant 或 linear")
    try:
        return fit_xas_global_components(
            datasets=[dataset.model_dump() for dataset in req.datasets],
            small_peak=req.small_peak.model_dump(),
            main_peaks=[peak.model_dump() for peak in req.main_peaks],
            fit_range=req.fit_range,
            main_background=req.main_background,
            ratio_numerator=req.ratio_numerator,
            ratio_denominator=req.ratio_denominator,
            max_nfev=req.max_nfev,
        )
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"全域擬合失敗：{exc}") from exc


@router.post("/global-fit-export/{export_format}")
def export_xas_global_fit(export_format: str, req: XasGlobalExportRequest):
    if export_format not in {"xlsx", "zip"}:
        raise HTTPException(status_code=400, detail="匯出格式僅支援 xlsx 或 zip")
    excel_bytes = _global_fit_excel(req.config, req.result)
    if export_format == "xlsx":
        return StreamingResponse(
            io.BytesIO(excel_bytes),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=xas_global_fit.xlsx"},
        )
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xas_global_fit_settings.json", json.dumps(req.config, ensure_ascii=False, indent=2))
        archive.writestr("xas_global_fit_results.json", json.dumps(req.result, ensure_ascii=False, indent=2))
        archive.writestr("xas_global_fit.xlsx", excel_bytes)
        for index, dataset in enumerate(req.result.get("datasets", []), start=1):
            components = dataset.get("components", [])
            peaks = dataset.get("peaks", [])
            columns: dict[str, Any] = {
                "Energy_eV": dataset.get("x", []),
                "Original": dataset.get("original", []),
                "Small_Gaussian": dataset.get("small_peak", {}).get("component", []),
                "Small_Local_Background": dataset.get("small_peak", {}).get("local_background", []),
                "Corrected": dataset.get("corrected", []),
                "Main_Background": dataset.get("background", []),
                "Total_Fit": dataset.get("total_fit", []),
                "Residual": dataset.get("residual", []),
            }
            for peak_index, component in enumerate(components):
                label = peaks[peak_index].get("label", f"Peak_{peak_index + 1}") if peak_index < len(peaks) else f"Peak_{peak_index + 1}"
                columns[f"Peak_{label}"] = component
            safe_name = re.sub(r"[^\w.-]+", "_", str(dataset.get("name") or f"sample_{index}"))
            archive.writestr(f"spectra/{index:02d}_{safe_name}.csv", pd.DataFrame(columns).to_csv(index=False))
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/zip",
        headers={"Content-Disposition": "attachment; filename=xas_global_fit_complete.zip"},
    )


class XasInitPeak(BaseModel):
    center: float
    fwhm: float
    amplitude: float
    label: Optional[str] = None
    lock_center: bool = True
    lock_fwhm: bool = True
    lock_area: bool = True
    center_min: Optional[float] = None
    center_max: Optional[float] = None
    fwhm_min: Optional[float] = None
    fwhm_max: Optional[float] = None
    amplitude_max: Optional[float] = None
    theoretical_center: Optional[float] = None


class XasFitRequest(BaseModel):
    x: List[float]
    y: List[float]
    peaks: List[XasInitPeak]
    profile: str = "voigt"
    maxfev: int = 6000
    peak_labels: Optional[List[str]] = None
    fit_range: Optional[List[float]] = None
    n_restarts: int = 1


class XasFitPeakRow(BaseModel):
    Peak_Name: str
    Center_eV: float
    FWHM_eV: float
    Area: float
    Height: float
    Area_pct: Optional[float] = None


class XasFitResponse(BaseModel):
    y_fit: List[float]
    y_individual: List[List[float]]
    residuals: List[float]
    peaks: List[XasFitPeakRow]
    r_squared: float = 0.0
    rmse: float = 0.0
    chi_red: Optional[float] = None


@router.post("/fit", response_model=XasFitResponse)
def fit_xas_peaks(req: XasFitRequest):
    x = np.array(req.x, dtype=float)
    y = np.array(req.y, dtype=float)

    if len(x) < 4 or not req.peaks:
        raise HTTPException(status_code=400, detail="資料或峰值參數不足")

    init_peaks = [
        {
            "center": pk.center,
            "fwhm": pk.fwhm,
            "amplitude": pk.amplitude,
            "lock_center": pk.lock_center,
            "lock_fwhm": pk.lock_fwhm,
            "lock_area": pk.lock_area,
            "center_min": pk.center_min,
            "center_max": pk.center_max,
            "fwhm_min": pk.fwhm_min,
            # Cap fwhm_max: if not provided, limit to 4× initial FWHM (max 15 eV)
            # to prevent peaks from growing unboundedly wide during XAS fitting.
            "fwhm_max": pk.fwhm_max if pk.fwhm_max is not None else min(max(pk.fwhm * 4.0, 2.0), 15.0),
            "amplitude_max": pk.amplitude_max,
            "theoretical_center": pk.theoretical_center,
            "label": pk.label,
        }
        for pk in req.peaks
    ]

    fit_range = req.fit_range
    if fit_range is None and req.peaks:
        centers = np.array([pk.center for pk in req.peaks], dtype=float)
        fwhms = np.array([max(pk.fwhm, 0.05) for pk in req.peaks], dtype=float)
        # Use median FWHM * 3 for padding, capped at 15 eV, to avoid including far background
        padding = min(max(5.0, float(np.median(fwhms)) * 3.0), 15.0)
        fit_lo = max(float(np.min(x)), float(np.min(centers)) - padding)
        fit_hi = min(float(np.max(x)), float(np.max(centers)) + padding)
        fit_range = [fit_lo, fit_hi]

    capped_maxfev = max(1000, min(int(req.maxfev), 8000))
    n_restarts = max(1, min(int(req.n_restarts), 10))
    rng = np.random.default_rng() if n_restarts > 1 else None

    best_result: dict | None = None
    best_ss_res = float("inf")
    for restart_idx in range(n_restarts):
        current_peaks = init_peaks if restart_idx == 0 else perturb_init_peaks(init_peaks, rng)
        try:
            candidate = fit_peaks(x, y, current_peaks, profile=req.profile,
                                  maxfev=capped_maxfev, fit_range=fit_range)
        except Exception:
            continue
        if not candidate.get("success", False):
            continue
        cand_ss = candidate.get("ss_res", float("inf"))
        if cand_ss < best_ss_res:
            best_result = candidate
            best_ss_res = cand_ss

    if best_result is None:
        raise HTTPException(status_code=422, detail="擬合失敗（全部嘗試均未收斂）")
    result = best_result

    raw_peaks = result.get("peaks", [])
    total_area = sum(abs(pk.get("area", 0)) for pk in raw_peaks)

    rows: list[XasFitPeakRow] = []
    for i, pk in enumerate(raw_peaks, 1):
        area = float(pk.get("area", 0))
        name = (
            req.peak_labels[i - 1]
            if req.peak_labels and i - 1 < len(req.peak_labels)
            else str(pk.get("label", f"Peak {i}"))
        )
        rows.append(XasFitPeakRow(
            Peak_Name=name,
            Center_eV=float(pk.get("center", 0)),
            FWHM_eV=float(pk.get("fwhm", 0)),
            Area=area,
            Height=float(pk.get("amplitude", 0)),
            Area_pct=round(100 * abs(area) / total_area, 2) if total_area > 0 else 0,
        ))

    def _to_list(arr):
        return arr.tolist() if hasattr(arr, "tolist") else list(arr)

    return XasFitResponse(
        y_fit=_to_list(result.get("y_fit", [])),
        y_individual=[_to_list(yi) for yi in result.get("y_individual", [])],
        residuals=_to_list(result.get("residuals", [])),
        peaks=rows,
        r_squared=float(result.get("r_squared", 0.0)),
        rmse=float(result.get("rmse", 0.0)),
        chi_red=result.get("chi_red"),
    )


# ── XAS sample database endpoints ─────────────────────────────────────────────

class XasSampleListItem(BaseModel):
    name: str
    description: str
    edges: List[str]


class XasEdgePeak(BaseModel):
    label: str
    energy_eV: float
    fwhm_eV: float
    meaning: str


class XasSampleEdgeResponse(BaseModel):
    sample: str
    edge: str
    energy_range: List[float]
    peaks: List[XasEdgePeak]


# ── Fit report (Excel) ────────────────────────────────────────────────────────

class FitReportPeak(BaseModel):
    name: str
    center: float
    fwhm: float
    area: float
    height: float
    area_pct: Optional[float] = None


class FitReportRequest(BaseModel):
    channel: str
    profile: str
    r2: float
    rmse: float
    chi_red: Optional[float] = None
    peaks: List[FitReportPeak]


@router.post("/fit-report")
def generate_fit_report(req: FitReportRequest):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl 未安裝")

    wb = Workbook()

    # ── Sheet 1: 擬合品質 ────────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "擬合品質"

    thin = Side(style="thin", color="CCCCCC")
    border_thin = Border(left=thin, right=thin, top=thin, bottom=thin)

    def _hdr_font(bold=True, size=11):
        return Font(name="Arial", bold=bold, size=size, color="FFFFFF")

    def _body_font(bold=False, size=10):
        return Font(name="Arial", bold=bold, size=size)

    def _fill(hex_color: str):
        return PatternFill(fill_type="solid", fgColor=hex_color)

    # Title
    ws1.merge_cells("A1:C1")
    t = ws1["A1"]
    t.value = "XAS 峰擬合分析報告"
    t.font = Font(name="Arial", bold=True, size=13, color="FFFFFF")
    t.fill = _fill("2D4A6B")
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 28

    # Section header
    ws1.merge_cells("A2:C2")
    s = ws1["A2"]
    s.value = "擬合參數"
    s.font = Font(name="Arial", bold=True, size=10, color="FFFFFF")
    s.fill = _fill("4A7CA8")
    s.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws1.row_dimensions[2].height = 20

    params_rows = [
        ("通道", req.channel.upper()),
        ("峰形", req.profile.upper()),
        ("峰數量", len(req.peaks)),
    ]
    for r_idx, (label, val) in enumerate(params_rows, start=3):
        ws1.cell(r_idx, 1, label).font = _body_font(bold=True)
        ws1.cell(r_idx, 1).fill = _fill("EDF3F9")
        ws1.cell(r_idx, 1).border = border_thin
        ws1.cell(r_idx, 2, val).font = _body_font()
        ws1.cell(r_idx, 2).border = border_thin
        ws1.cell(r_idx, 3).border = border_thin

    # Section header: 擬合品質
    next_r = 3 + len(params_rows)
    ws1.merge_cells(f"A{next_r}:C{next_r}")
    sq = ws1.cell(next_r, 1, "擬合品質指標")
    sq.font = Font(name="Arial", bold=True, size=10, color="FFFFFF")
    sq.fill = _fill("4A7CA8")
    sq.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws1.row_dimensions[next_r].height = 20

    # R² with conditional color
    r2_val = req.r2
    if r2_val >= 0.99:
        r2_fill = "D4EDDA"; r2_txt = "228B22"
    elif r2_val >= 0.97:
        r2_fill = "D0E8F5"; r2_txt = "0066AA"
    elif r2_val >= 0.90:
        r2_fill = "FFF3CD"; r2_txt = "856404"
    else:
        r2_fill = "F8D7DA"; r2_txt = "842029"

    quality_rows: list[tuple] = [
        ("R²", f"{r2_val:.6f}", r2_fill, r2_txt),
        ("RMSE", f"{req.rmse:.6f}", "FFFFFF", "000000"),
    ]
    if req.chi_red is not None:
        quality_rows.append(("χ²ᵣ (Reduced χ²)", f"{req.chi_red:.6f}", "FFFFFF", "000000"))

    for q_idx, (label, val, bg, fg) in enumerate(quality_rows, start=next_r + 1):
        c_label = ws1.cell(q_idx, 1, label)
        c_label.font = Font(name="Arial", bold=True, size=10, color="000000")
        c_label.fill = _fill("EDF3F9")
        c_label.border = border_thin
        c_val = ws1.cell(q_idx, 2, val)
        c_val.font = Font(name="Arial", bold=(bg != "FFFFFF"), size=10, color=fg)
        c_val.fill = _fill(bg)
        c_val.border = border_thin
        ws1.cell(q_idx, 3).border = border_thin

    ws1.column_dimensions["A"].width = 24
    ws1.column_dimensions["B"].width = 18
    ws1.column_dimensions["C"].width = 4

    # ── Sheet 2: 峰參數 ──────────────────────────────────────────────────
    ws2 = wb.create_sheet("峰參數")

    col_headers = ["峰名稱", "中心 (eV)", "FWHM (eV)", "面積", "高度", "面積 %"]
    col_widths = [18, 14, 14, 14, 14, 10]
    hdr_fill = _fill("2D4A6B")

    for c_idx, (hdr, w) in enumerate(zip(col_headers, col_widths), start=1):
        cell = ws2.cell(1, c_idx, hdr)
        cell.font = _hdr_font()
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border_thin
        ws2.column_dimensions[get_column_letter(c_idx)].width = w
    ws2.row_dimensions[1].height = 22

    for r_idx, pk in enumerate(req.peaks, start=2):
        row_bg = "F7FAFD" if r_idx % 2 == 0 else "FFFFFF"
        values = [pk.name, pk.center, pk.fwhm, pk.area, pk.height, pk.area_pct]
        for c_idx, val in enumerate(values, start=1):
            cell = ws2.cell(r_idx, c_idx, val if val is not None else "—")
            cell.font = _body_font()
            cell.fill = _fill(row_bg)
            cell.border = border_thin
            if c_idx > 1 and isinstance(val, float):
                cell.number_format = "0.0000" if c_idx <= 5 else "0.00"
            if c_idx == 1:
                cell.alignment = Alignment(horizontal="left", indent=1)
            else:
                cell.alignment = Alignment(horizontal="right")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=xas_fit_report.xlsx"},
    )


@router.get("/samples", response_model=List[XasSampleListItem])
def get_xas_samples():
    return [XasSampleListItem(**s) for s in list_samples()]


@router.get("/sample-peaks/{sample_name}/{edge_name:path}", response_model=XasSampleEdgeResponse)
def get_xas_sample_peaks(sample_name: str, edge_name: str):
    data = get_sample_edge_peaks(sample_name, edge_name)
    if data is None:
        raise HTTPException(status_code=404, detail=f"找不到樣品 '{sample_name}' 或邊 '{edge_name}'")
    return XasSampleEdgeResponse(
        sample=data["sample"],
        edge=data["edge"],
        energy_range=data["energy_range"],
        peaks=[XasEdgePeak(**p) for p in data["peaks"]],
    )
